"""Build the user guide from guide_content.py:
    help.html                      (published with the site, opened from Help in the sidebar)
    guide/Project_Hub_User_Guide.pdf
    guide/Project_Hub_User_Guide.docx

Run the local server on :8765 first (the PDF is printed from it), then:
    py -3 tools/build_guide.py
"""
import html
import re

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for easy landscape switch)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

from guide_content import BLOCKS, SUBTITLE, TITLE, VERSION
from hub import BASE, ROOT

SRC = ROOT / 'guide' / 'img'
WEB = ROOT / 'guide' / 'web'
OUT_PDF = ROOT / 'guide' / 'Project_Hub_User_Guide.pdf'
OUT_DOCX = ROOT / 'guide' / 'Project_Hub_User_Guide.docx'
ACCENT = RGBColor(0x1D, 0x5B, 0x79)
INK = RGBColor(0x18, 0x22, 0x2A)
MUTED = RGBColor(0x5D, 0x6A, 0x73)


def web_images():
    WEB.mkdir(parents=True, exist_ok=True)
    for src in sorted(SRC.glob('*.png')):
        im = Image.open(src).convert('RGB')
        w = 1600
        im = im.resize((w, round(im.height * w / im.width)), Image.LANCZOS)
        im.save(WEB / src.name, optimize=True)


def inline_html(text):
    return re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', html.escape(text))


def slug(text):
    return re.sub(r'[^a-z0-9]+', '-', text.lower()).strip('-')


# ------------------------------------------------------------------ HTML

def build_html():
    toc, body = [], []
    n1 = 0
    n2 = 0
    for b in BLOCKS:
        kind = b[0]
        if kind == 'h1':
            n1 += 1
            n2 = 0
            toc.append(f'<li><a href="#{b[2]}">{n1}. {html.escape(b[1])}</a><ol>')
            body.append(f'{"</section>" if n1 > 1 else ""}<section class="chapter"><h1 id="{b[2]}"><span>{n1}</span>{html.escape(b[1])}</h1>')
        elif kind == 'h2':
            n2 += 1
            a = f'{BLOCKS_ANCHOR(n1)}-{slug(b[1])}'
            toc[-1] += f'<li><a href="#{a}">{html.escape(b[1])}</a></li>'
            body.append(f'<h2 id="{a}">{n1}.{n2} {html.escape(b[1])}</h2>')
        elif kind == 'p':
            body.append(f'<p>{inline_html(b[1])}</p>')
        elif kind == 'steps':
            body.append('<ol class="steps">' + ''.join(f'<li>{inline_html(s)}</li>' for s in b[1]) + '</ol>')
        elif kind == 'bullets':
            body.append('<ul>' + ''.join(f'<li>{inline_html(s)}</li>' for s in b[1]) + '</ul>')
        elif kind == 'img':
            body.append(f'<figure><img src="guide/web/{b[1]}.png" alt="{html.escape(b[2])}" loading="lazy"><figcaption>{html.escape(b[2])}</figcaption></figure>')
        elif kind in ('tip', 'warn'):
            body.append(f'<div class="callout {kind}"><b>{"Tip" if kind == "tip" else "Note"}</b> {inline_html(b[1])}</div>')
        elif kind == 'table':
            head = ''.join(f'<th>{html.escape(h)}</th>' for h in b[1])
            rows = ''.join('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in r) + '</tr>' for r in b[2])
            body.append(f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>')
    toc = [t + '</ol></li>' for t in toc]
    body.append('</section>')
    page = f'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{TITLE}</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans+Condensed:wght@600&family=IBM+Plex+Sans:wght@400;500;600&display=swap">
<style>
:root {{ --bg:#f2f4f3; --surface:#fff; --ink:#18222a; --ink-2:#47545e; --ink-3:#75818a; --line:#e0e4e2; --accent:#1d5b79; --soft:#e2edf2; --warn:#a96300; --warn-soft:#fbeed6; --side:#17232c; color-scheme:light; }}
@media (prefers-color-scheme: dark) {{ :root {{ --bg:#11171c; --surface:#182128; --ink:#e3e9ed; --ink-2:#a8b4bd; --ink-3:#7b8892; --line:#2a3640; --accent:#72b5d4; --soft:#1b3341; --warn:#e3aa4c; --warn-soft:#34280f; --side:#0c1216; color-scheme:dark; }} }}
* {{ box-sizing:border-box }}
body {{ margin:0; background:var(--bg); color:var(--ink); font:15px/1.65 "IBM Plex Sans","Segoe UI",system-ui,sans-serif; }}
.wrap {{ display:grid; grid-template-columns:280px minmax(0,1fr); min-height:100vh; }}
nav {{ background:var(--side); color:#b9c6cf; padding:24px 18px; position:sticky; top:0; height:100vh; overflow-y:auto; font-size:13.5px; }}
nav b {{ display:block; color:#fff; font:600 17px "IBM Plex Sans Condensed",sans-serif; margin-bottom:4px; }}
nav .dl {{ display:flex; gap:8px; margin:14px 0 18px; flex-wrap:wrap; }}
nav .dl a {{ color:#fff; border:1px solid #3a4c58; border-radius:6px; padding:5px 10px; text-decoration:none; font-weight:500; }}
nav .dl a:hover {{ background:#22313c; }}
nav ol {{ list-style:none; margin:0; padding:0; }}
nav > ol > li {{ margin:10px 0; }}
nav > ol > li > a {{ color:#fff; font-weight:600; }}
nav ol ol {{ padding-left:12px; margin-top:4px; }}
nav a {{ color:#b9c6cf; text-decoration:none; display:block; padding:2px 0; }}
nav a:hover {{ color:#fff; }}
main {{ min-width:0; padding:40px max(16px,4vw) 80px; }}
.doc {{ max-width:900px; margin:0 auto; }}
.cover {{ background:var(--surface); border:1px solid var(--line); border-radius:12px; padding:36px; margin-bottom:28px; }}
.cover h1 {{ font:600 34px/1.15 "IBM Plex Sans Condensed",sans-serif; margin:0 0 6px; }}
.cover p {{ margin:0; color:var(--ink-2); }}
.cover img {{ width:100%; border-radius:8px; border:1px solid var(--line); margin-top:22px; }}
.chapter {{ background:var(--surface); border:1px solid var(--line); border-radius:12px; padding:30px 36px; margin-bottom:28px; }}
h1 {{ font:600 26px/1.2 "IBM Plex Sans Condensed",sans-serif; margin:0 0 12px; display:flex; align-items:center; gap:12px; }}
h1 span {{ display:inline-grid; place-items:center; width:36px; height:36px; border-radius:8px; background:var(--accent); color:#fff; font-size:18px; flex:none; }}
h2 {{ font:600 19px/1.3 "IBM Plex Sans Condensed",sans-serif; margin:30px 0 8px; color:var(--accent); }}
p, li {{ max-width:72ch; }}
ol.steps {{ counter-reset:s; list-style:none; padding:0; }}
ol.steps li {{ counter-increment:s; position:relative; padding-left:38px; margin:8px 0; }}
ol.steps li::before {{ content:counter(s); position:absolute; left:0; top:1px; width:24px; height:24px; border-radius:50%; background:var(--soft); color:var(--accent); font-weight:600; font-size:13px; display:grid; place-items:center; }}
figure {{ margin:18px 0 22px; }}
figure img {{ width:100%; display:block; border:1px solid var(--line); border-radius:8px; box-shadow:0 4px 18px rgba(16,28,38,.08); }}
figcaption {{ font-size:13px; color:var(--ink-3); margin-top:6px; }}
.callout {{ border-radius:8px; padding:10px 14px; margin:14px 0; background:var(--soft); }}
.callout b {{ color:var(--accent); margin-right:4px; }}
.callout.warn {{ background:var(--warn-soft); }}
.callout.warn b {{ color:var(--warn); }}
.tw {{ overflow-x:auto; margin:12px 0 18px; }}
table {{ border-collapse:collapse; width:100%; font-size:14px; }}
th {{ text-align:left; background:var(--soft); padding:8px 10px; font-weight:600; }}
td {{ border-bottom:1px solid var(--line); padding:8px 10px; vertical-align:top; }}
td:first-child {{ font-weight:500; white-space:nowrap; }}
@media (max-width:860px) {{ .wrap {{ grid-template-columns:1fr; }} nav {{ position:static; height:auto; }} .chapter, .cover {{ padding:22px 18px; }} td:first-child {{ white-space:normal; }} }}
@media print {{
  @page {{ size:A4; margin:16mm 15mm 18mm; }}
  body {{ background:#fff; font-size:10.5pt; color:#18222a; }}
  .wrap {{ display:block; }} nav {{ display:none; }} main {{ padding:0; }}
  .cover {{ border:0; padding:60mm 0 0; min-height:250mm; page-break-after:always; }}
  .cover h1 {{ font-size:32pt; }}
  .chapter {{ border:0; padding:0; margin:0; page-break-before:always; }}
  figure {{ page-break-inside:avoid; }} figure img {{ box-shadow:none; }}
  h2, h1 {{ page-break-after:avoid; }} tr, .callout {{ page-break-inside:avoid; }}
  .toc-print {{ display:block !important; page-break-after:always; }}
}}
.toc-print {{ display:none; }}
.toc-print h1 {{ display:block; }}
.toc-print ol {{ font-size:11pt; line-height:1.8; list-style:none; padding-left:0; }}
.toc-print ol ol {{ padding-left:16px; }}
.toc-print ol ol {{ font-size:10pt; line-height:1.6; color:#47545e; }}
</style>
</head>
<body>
<div class="wrap">
<nav aria-label="Contents"><b>{TITLE}</b><span>{VERSION}</span>
<div class="dl"><a href="guide/Project_Hub_User_Guide.pdf" download>Download PDF</a><a href="guide/Project_Hub_User_Guide.docx" download>Download Word</a><a href="./">Open Project Hub</a></div>
<ol>{"".join(toc)}</ol></nav>
<main><div class="doc">
<header class="cover"><h1>{TITLE}</h1><p>{SUBTITLE}</p><p style="margin-top:6px">{VERSION}</p><img src="guide/web/02_dashboard.png" alt="Project Hub dashboard"></header>
<section class="toc-print"><h1>Contents</h1><ol>{"".join(toc)}</ol></section>
{"".join(body)}
</div></main>
</div>
</body>
</html>'''
    (ROOT / 'help.html').write_text(page, encoding='utf-8')
    print('help.html written')


def BLOCKS_ANCHOR(n1):
    return [b[2] for b in BLOCKS if b[0] == 'h1'][n1 - 1]


# ------------------------------------------------------------------ PDF

def build_pdf():
    with sync_playwright() as p:
        browser = p.chromium.launch(channel='msedge', headless=True)
        page = browser.new_page()
        page.goto(BASE + 'help.html', wait_until='networkidle')
        page.emulate_media(media='print')
        page.pdf(
            path=str(OUT_PDF), format='A4', print_background=True, display_header_footer=True,
            header_template='<span></span>',
            footer_template=f'<div style="font:8px Segoe UI,sans-serif;color:#75818a;width:100%;padding:0 15mm;display:flex;justify-content:space-between"><span>{TITLE}</span><span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>',
            margin={'top': '16mm', 'bottom': '18mm', 'left': '15mm', 'right': '15mm'},
        )
        browser.close()
    print('PDF written', OUT_PDF.stat().st_size // 1024, 'KB')


# ------------------------------------------------------------------ Word

def runs(par, text, size=None, color=None):
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = i % 2 == 1
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return par


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def field(par, instr):
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, ''), ('end', None)):
        r = par.add_run()
        if kind:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), kind)
            r._r.append(el)
        elif text is not None and text:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
            r._r.append(el)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    for lvl, size in ((1, 20), (2, 14)):
        hs = doc.styles[f'Heading {lvl}']
        hs.font.name = 'Calibri'
        hs.font.size = Pt(size)
        hs.font.color.rgb = ACCENT if lvl == 2 else INK
        hs.font.bold = True

    fp = sec.footer.paragraphs[0]
    runs(fp, TITLE + '    ·    page ', size=8, color=MUTED)
    field(fp, 'PAGE')
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    runs(t, TITLE, size=30, color=INK).runs[0].bold = True
    runs(doc.add_paragraph(), SUBTITLE, size=13, color=MUTED)
    runs(doc.add_paragraph(), VERSION, size=11, color=MUTED)
    doc.add_paragraph()
    doc.add_picture(str(WEB / '02_dashboard.png'), width=Cm(17))
    doc.add_page_break()

    # Contents (Word fills it in when the file is opened and fields are updated)
    doc.add_heading('Contents', level=1)
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    runs(doc.add_paragraph(), 'If the contents list is empty, right-click here and choose Update Field.', size=9, color=MUTED)

    n1 = n2 = 0
    for b in BLOCKS:
        kind = b[0]
        if kind == 'h1':
            n1 += 1
            n2 = 0
            doc.add_page_break()
            doc.add_heading(f'{n1}. {b[1]}', level=1)
        elif kind == 'h2':
            n2 += 1
            doc.add_heading(f'{n1}.{n2} {b[1]}', level=2)
        elif kind == 'p':
            runs(doc.add_paragraph(), b[1])
        elif kind == 'steps':
            for i, s in enumerate(b[1], 1):
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.first_line_indent = Cm(-0.6)
                par.paragraph_format.space_after = Pt(3)
                r = par.add_run(f'{i}.  ')
                r.bold = True
                r.font.color.rgb = ACCENT
                runs(par, s)
        elif kind == 'bullets':
            for s in b[1]:
                runs(doc.add_paragraph(style='List Bullet'), s)
        elif kind == 'img':
            doc.add_picture(str(WEB / f'{b[1]}.png'), width=Cm(17))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = runs(doc.add_paragraph(), b[2], size=9, color=MUTED)
            cap.runs[0].italic = True
        elif kind in ('tip', 'warn'):
            tb = doc.add_table(rows=1, cols=1)
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tb.rows[0].cells[0]
            shade(c, 'E2EDF2' if kind == 'tip' else 'FBEED6')
            par = c.paragraphs[0]
            r = par.add_run(('Tip  ' if kind == 'tip' else 'Note  '))
            r.bold = True
            r.font.color.rgb = ACCENT if kind == 'tip' else RGBColor(0xA9, 0x63, 0x00)
            runs(par, b[1])
            doc.add_paragraph()
        elif kind == 'table':
            tb = doc.add_table(rows=1, cols=len(b[1]))
            tb.style = 'Table Grid'
            for i, h in enumerate(b[1]):
                c = tb.rows[0].cells[i]
                shade(c, 'E2EDF2')
                c.paragraphs[0].add_run(h).bold = True
            for row in b[2]:
                cells = tb.add_row().cells
                for i, v in enumerate(row):
                    runs(cells[i].paragraphs[0], v, size=9.5)
            for row in tb.rows:
                row.cells[0].width = Cm(4.6)
                for c in row.cells[1:]:
                    c.width = Cm(12.4 / max(1, len(row.cells) - 1))
            doc.add_paragraph()
    doc.save(str(OUT_DOCX))
    print('Word written', OUT_DOCX.stat().st_size // 1024, 'KB')


if __name__ == '__main__':
    web_images()
    build_html()
    build_pdf()
    build_docx()
